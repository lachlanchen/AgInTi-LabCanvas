#!/usr/bin/env python3
"""Editable PowerPoint version of the paper figures: image + native label boxes + leader lines.

Run with a Python that has python-pptx (the miniconda base python3 does):
  python3 make_paper_figure_pptx.py [--design-dir DIR] [--figures figA2 figF ...]
Output: artifacts/paper_figure/paper_figures_labelled.pptx (one slide per figure, 4:3 to match the
2400 x 1800 renders). Every label is a real text box with a straight connector, so text, font, box
style and positions can be edited in PowerPoint/Keynote.
"""
from __future__ import annotations

import argparse
import json
import math
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

SLIDE_W, SLIDE_H = 13.333, 10.0          # inches, 4:3 like the renders
IMG_W_PX, IMG_H_PX = 2400, 1800
FONT = "Arial"
FONT_SIZE = Pt(14)
TEXT_RGB = RGBColor(0x22, 0x22, 0x22)
LINE_RGB = RGBColor(0x55, 0x55, 0x55)
BOX_LINE_RGB = RGBColor(0x8C, 0x8C, 0x8C)
BOX_H = 0.44

# clean label wording for the paper; provenance goes to the slide notes
LABELS = {
    "holder": "EVK5 camera holder",
    "grating": "Diffraction grating",
    "sensor_board": "Sensor board (EBS)",
    "light_path": "Light path",
    "upper_shelf_window": "Shelf window",
    "incubator": "Incubator",
    "slider": "Linear stage carriage",
    "led_holder": "LED holder with Lumileds PCB",
    "lumileds_pcb": "Lumileds LED PCB",
}
# per figure: label -> (text centre x fraction, y fraction[, anchor key])
FIGS = {
    "figA2_incubator_overview_clean_slider": {
        "slider": (0.83, 0.13), "led_holder": (0.83, 0.25),
        "grating": (0.84, 0.55, "grating_edge"), "holder": (0.84, 0.80, "holder_top"),
        "sensor_board": (0.16, 0.86, "sensor_board_left"), "incubator": (0.12, 0.30)},
    "figF_front_elevation_alignment": {
        "slider": (0.80, 0.10), "led_holder": (0.80, 0.21), "grating": (0.20, 0.50, "grating_edge"),
        "holder": (0.80, 0.72, "holder_top"), "sensor_board": (0.18, 0.80, "sensor_board_left")},
    "figA_incubator_overview": {
        "holder": (0.84, 0.80, "holder_top"), "grating": (0.84, 0.55, "grating_edge"),
        "sensor_board": (0.16, 0.86, "sensor_board_left"), "light_path": (0.64, 0.12, "light_path_high"), "incubator": (0.12, 0.30)},
    "figB_optical_stack": {
        "holder": (0.82, 0.80, "holder_top"), "grating": (0.20, 0.22, "grating_edge"),
        "sensor_board": (0.14, 0.72, "sensor_board_left"), "light_path": (0.64, 0.08, "light_path_low"),
        "upper_shelf_window": (0.84, 0.38, "upper_shelf_window_right")},
    "figC_holder_under_grating_closeup": {
        "holder": (0.82, 0.84, "holder_top"), "grating": (0.18, 0.18, "grating_edge"),
        "sensor_board": (0.14, 0.86, "sensor_board_left"), "light_path": (0.64, 0.07, "light_path_low")},
    "figE_section_through_light_path": {
        "holder": (0.84, 0.60, "holder_top"), "grating": (0.20, 0.34, "grating_edge"),
        "sensor_board": (0.16, 0.76, "sensor_board_left"), "light_path": (0.62, 0.08, "light_path_low"),
        "upper_shelf_window": (0.84, 0.44, "upper_shelf_window_right")},
}
NOTES = {
    "figA2_incubator_overview_clean_slider": "Incubator thinner (Shapr3D export) with the EVK 5 holder group translated onto the sensor board (dy +220, dz +129 mm). The linear stage carriage, LED holder and Lumileds PCB are a proposed illustration layer (cad/designs/incubator_thinner_evk5_holder_on_sensor_board, manifest.json > figure_proposal). Vertical rails and the light-path marker are hidden in this view.",
    "figF_front_elevation_alignment": "Orthographic front elevation: carriage, LED holder, grating window and EVK5 holder share the light-path X = 141.25 mm.",
}


def px_to_in(x_px: float, y_px: float) -> tuple[float, float]:
    return x_px / IMG_W_PX * SLIDE_W, y_px / IMG_H_PX * SLIDE_H


def box_size(text: str) -> tuple[float, float]:
    return max(1.4, 0.115 * len(text) + 0.45), BOX_H


def edge_point(cx: float, cy: float, w: float, h: float, tx: float, ty: float) -> tuple[float, float]:
    """Point on the box border along the ray from the box centre to the target."""
    dx, dy = tx - cx, ty - cy
    if abs(dx) < 1e-9 and abs(dy) < 1e-9:
        return cx, cy
    sx = (w / 2) / abs(dx) if dx else math.inf
    sy = (h / 2) / abs(dy) if dy else math.inf
    s = min(sx, sy)
    return cx + dx * s, cy + dy * s


def add_label(slide, text: str, cx: float, cy: float, ax: float, ay: float) -> None:
    w, h = box_size(text)
    left, top = cx - w / 2, cy - h / 2
    # leader line from box edge to anchor, drawn first so the box sits on top
    ex, ey = edge_point(cx, cy, w, h, ax, ay)
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(ex), Inches(ey), Inches(ax), Inches(ay))
    line.line.color.rgb = LINE_RGB
    line.line.width = Pt(1.0)
    line.name = f"leader: {text}"
    # anchor dot
    d = 0.10
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(ax - d / 2), Inches(ay - d / 2), Inches(d), Inches(d))
    dot.fill.solid()
    dot.fill.fore_color.rgb = LINE_RGB
    dot.line.fill.background()
    dot.shadow.inherit = False
    dot.name = f"anchor: {text}"
    # label box
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h))
    box.adjustments[0] = 0.18
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    box.line.color.rgb = BOX_LINE_RGB
    box.line.width = Pt(0.75)
    box.shadow.inherit = False
    box.name = f"label: {text}"
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.10)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = FONT_SIZE
    run.font.bold = False
    run.font.color.rgb = TEXT_RGB


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--design-dir", type=Path, default=Path(__file__).resolve().parent)
    ap.add_argument("--figures", nargs="*", default=list(FIGS))
    ap.add_argument("--out", type=Path)
    args = ap.parse_args()
    figdir = args.design_dir / "artifacts" / "paper_figure"
    out = args.out or figdir / "paper_figures_labelled.pptx"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    blank = prs.slide_layouts[6]
    made = 0
    for name in args.figures:
        key = next((k for k in FIGS if k.startswith(name)), None)
        img = figdir / f"{key}.png" if key else None
        if not key or not img.exists():
            print("skip", name)
            continue
        anchors = json.loads((figdir / f"{key}_labels.json").read_text())
        slide = prs.slides.add_slide(blank)
        pic = slide.shapes.add_picture(str(img), 0, 0, width=Inches(SLIDE_W), height=Inches(SLIDE_H))
        pic.name = key
        for label_key, spec in FIGS[key].items():
            fx, fy = spec[0], spec[1]
            akey = spec[2] if len(spec) > 2 else label_key
            if akey not in anchors:
                continue
            ax_px, ay_px = anchors[akey]["x_px"], anchors[akey]["y_px"]
            if not (0 <= ax_px <= IMG_W_PX and 0 <= ay_px <= IMG_H_PX):
                continue
            ax, ay = px_to_in(ax_px, ay_px)
            add_label(slide, LABELS[label_key], fx * SLIDE_W, fy * SLIDE_H, ax, ay)
        if key in NOTES:
            slide.notes_slide.notes_text_frame.text = NOTES[key]
        made += 1
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"wrote {out} ({made} slides)")


if __name__ == "__main__":
    main()
